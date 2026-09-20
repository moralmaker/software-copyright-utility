#!/usr/bin/env python3
"""在官方《计算机软件著作权登记申请表》模板上原位填写，生成正式登记表 DOCX。

读取已确认的 草稿/申请表信息.md，把值写入模板对应的单元格，保持模板原有
版式、边框、字体与分区样式不变：
- 文本字段：写入对应空白值单元格；
- 选项字段（分类/软件作品说明/发表状态/开发方式/权利取得方式/权利范围/编程语言/
  软件技术类型/申请方式）：按模板复选框样式打勾，选中项的 <w:sym> 设为勾选框(☑)、
  其余清回空框(□)；模板里用字面 ● 表示的复选框（如“由代理人申请”）先转成标准空框(□)。
- 带“（限 50 个字符）”提示的字段：在提示上方写入实际值；
- 著作权人：写入第一数据行（姓名或名称/类别/证件类型/证件号码/国籍/省份城市）。

模板默认取 skill 内 assets/计算机软件著作权登记申请表-模板.docx，可用 --template 覆盖
（例如项目里 doc/sc/计算机软件著作权登记申请表.docx）。
"""

from __future__ import annotations

import argparse
import copy
import re
import shutil
from pathlib import Path
from typing import Any, Optional

from common import ensure_dir

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    DOCX_AVAILABLE = True
except Exception:
    DOCX_AVAILABLE = False

FONT = "SimSun"
DEFAULT_SIZE = 10.5

TECH_TYPES = [
    "APP", "游戏软件", "教育软件", "金融软件", "医疗软件", "地理信息软件",
    "云计算软件", "信息安全软件", "大数据软件", "人工智能软件", "VR软件",
    "5G软件", "小程序", "物联网软件", "智慧城市软件",
]


def parse_application_fields(md_path: Path) -> dict[str, str]:
    fields: dict[str, str] = {}
    if not md_path.exists():
        return fields
    for line in md_path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if stripped.startswith("➤") and "：" in stripped:
            key, _, value = stripped[1:].partition("：")
            fields[key.strip()] = value.strip()
    return fields


def format_date(value: str) -> str:
    value = (value or "").strip()
    m = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", value)
    if m:
        return f"{m.group(1)} 年 {m.group(2).zfill(2)} 月 {m.group(3).zfill(2)} 日"
    return value


def parse_copyright_holder(raw: str) -> tuple[str, str, str, str]:
    name = (raw or "").strip()
    category, cert_type, cert_no = "", "", ""
    m = re.match(r"^(.*)[（(]([^）)]*)[）)]$", name)
    if m and any(k in m.group(2) for k in ("证件类型", "证件号", "法人", "自然人")):
        name = m.group(1).strip()
        for part in re.split(r"[，,、]", m.group(2)):
            part = part.strip()
            if not part:
                continue
            if part.startswith("证件类型"):
                cert_type = re.split(r"[：:]", part, 1)[-1].strip()
            elif part.startswith("证件号"):
                cert_no = re.split(r"[：:]", part, 1)[-1].strip()
            elif "法人" in part or "自然人" in part:
                category = "法人" if "法人" in part else "自然人"
    if not (category or cert_type or cert_no):
        name = (raw or "").strip()
    return name, category, cert_type, cert_no


def region_from_name(name: str) -> str:
    m = re.search(r"[（(]([^（）()]{2,4})[）)]", name or "")
    return m.group(1) if m else ""


def _nested_tables(cell) -> list:
    from docx.table import Table
    return [Table(t, cell) for t in cell._tc.findall(qn("w:tbl"))]


def _dedup(row) -> list:
    out, seen = [], set()
    for c in row.cells:
        if c._tc in seen:
            continue
        seen.add(c._tc)
        out.append(c)
    return out


def _first_font(cell):
    for p in cell.paragraphs:
        for r in p.runs:
            if r.text.strip():
                return (r.font.name or FONT, r.font.size.pt if r.font.size else DEFAULT_SIZE)
    return None


def _set_run(run, font: tuple[str, float]):
    name, size = font
    run.font.name = name
    run.font.size = Pt(size)
    try:
        run.font.color.rgb = run.font.color.rgb  # keep template color (black)
    except Exception:
        pass
    try:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    except Exception:
        pass


def set_cell(cell, text: str, font: Optional[tuple[str, float]] = None) -> None:
    font = font or (FONT, DEFAULT_SIZE)
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    lines = (text or "").split("\n")
    if not lines:
        lines = [""]
    for line in lines:
        p = cell.add_paragraph()
        run = p.add_run(line if line else " ")
        _set_run(run, font)


def value_font(label_cell, option_cell=None) -> tuple[str, float]:
    for c in (label_cell, option_cell):
        f = _first_font(c)
        if f:
            return f
    return (FONT, DEFAULT_SIZE)


WINGDINGS_EMPTY = "00A3"    # Wingdings 2 空框 □
WINGDINGS_CHECKED = "0052"  # Wingdings 2 勾选框 ☑


def _option_segments(p_elem) -> list:
    """返回段落内 [(sym_elem, label_text), ...]。

    每个 <w:sym> 是一个选项的复选框，其 label 为到下一个 <w:sym> 之间的文本。
    段落首个 <w:sym> 之前的文本对应 sym=None（忽略）。
    """
    sym_tag = qn("w:sym")
    r_tag = qn("w:r")
    t_tag = qn("w:t")
    segs: list = []
    cur_sym = None
    buf: list = []
    for r in p_elem.iterchildren(r_tag):
        sym = r.find(sym_tag)
        if sym is not None:
            segs.append((cur_sym, "".join(buf)))
            cur_sym = sym
            buf = []
        else:
            for t in r.findall(t_tag):
                buf.append(t.text or "")
    segs.append((cur_sym, "".join(buf)))
    return segs


def mark_option(cell, tokens, font=None) -> None:
    """按模板复选框样式打勾：选中项的 <w:sym> 改为勾选框(0052)，其余清回空框(00A3)。

    tokens 为需勾选的选项文本列表（多选则传多个）。词边界匹配，避免 Java 命中 JavaScript。
    """
    token_list = [t for t in (tokens or []) if t]
    pats = [
        re.compile(r"(?<![A-Za-z0-9])" + re.escape(t) + r"(?![A-Za-z0-9])")
        for t in token_list
    ]

    def wanted(label: str) -> bool:
        s = label.strip()
        return any(p.search(s) for p in pats)

    for p in cell.paragraphs:
        for sym, label in _option_segments(p._p):
            if sym is None:
                continue
            sym.set(qn("w:char"), WINGDINGS_CHECKED if wanted(label) else WINGDINGS_EMPTY)


def _normalize_checkbox_bullets(cell) -> None:
    """把模板里用字面 ●/○ 文本 run 表示的复选框，替换成标准 <w:sym>(Wingdings 2) 空框 run。

    例如申请人格“由代理人申请”前模板用的是字面 ●，这里转成真正的 □ 空框，
    使其与“由著作权人申请”一样具备可勾选的复选框样式。
    """
    template_run = None
    for p in cell.paragraphs:
        for r in p._p.iterchildren(qn("w:r")):
            if r.find(qn("w:sym")) is not None:
                template_run = r
                break
        if template_run is not None:
            break
    if template_run is None:
        return
    for p in cell.paragraphs:
        for r in list(p._p.iterchildren(qn("w:r"))):
            if r.find(qn("w:sym")) is not None:
                continue
            txt = "".join(t.text or "" for t in r.findall(qn("w:t")))
            if txt and all(ch in "●○" for ch in txt):
                new_run = copy.deepcopy(template_run)
                for sym in new_run.findall(qn("w:sym")):
                    sym.set(qn("w:char"), WINGDINGS_EMPTY)
                r.addprevious(new_run)
                r.getparent().remove(r)


def _fill_first_publish_date(cell, date_str: str) -> None:
    """把首次发表日期填入 年/月/日 占位 run，保留同段的已发表/未发表复选框符号。"""
    m = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", date_str or "")
    if not m:
        return
    y, mo, day = m.group(1), m.group(2).zfill(2), m.group(3).zfill(2)
    for p in cell.paragraphs:
        for r in list(p._p.iterchildren(qn("w:r"))):
            t = r.find(qn("w:t"))
            if t is None or t.text is None:
                continue
            s = t.text.strip()
            if s == "年":
                t.text = f"{y} 年"
            elif s == "月":
                t.text = f"{mo} 月"
            elif s == "日" and "年" in p.text:
                t.text = f"{day} 日"


def add_value_above_hint(cell, value: str, font) -> None:
    """在“（...限50个字符）”提示上方写入实际值，保留提示。"""
    hint = cell.text.strip()
    set_cell(cell, (value + "\n" + hint) if hint else value, font)


def build_form(software_name: str, fields: dict[str, str], template_path: Path, out_path: Path) -> None:
    shutil.copyfile(template_path, out_path)
    doc = Document(str(out_path))

    g = lambda key: fields.get(key, "")
    t1 = doc.tables[1]
    t3 = doc.tables[3]

    soft = _nested_tables(t1.rows[0].cells[1])[0]        # 软件基本信息
    date_tbl = _nested_tables(t1.rows[1].cells[1])[0]    # 开发完成日期/发表状态/开发方式
    holder = _nested_tables(t1.rows[2].cells[1])[0]      # 著作权人
    rights = _nested_tables(t3.rows[0].cells[1])[0]      # 权利说明
    tech = _nested_tables(t3.rows[2].cells[1])[0]        # 软件功能和技术特点
    applicant = _nested_tables(t3.rows[3].cells[1])[0]   # 申请人信息

    # ---- 软件基本信息 ----
    s0, s1, s2 = _dedup(soft.rows[0]), _dedup(soft.rows[1]), _dedup(soft.rows[2])
    # 软件名称 s0C1, 版本号 s0C3
    set_cell(s0[1], g("软件全称"), value_font(s0[0]))
    set_cell(s0[3], g("版本号"), value_font(s0[2]))
    # 软件简称 s1C1, 分类 s1C3
    set_cell(s1[1], g("软件简称") or "无", value_font(s1[0]))
    mark_option(s1[3], [g("软件分类")], value_font(s1[2], s1[3]))
    # 软件作品说明 s2C1(原创)
    if g("软件说明"):
        mark_option(s2[1], [g("软件说明")], value_font(s2[0], s2[1]))

    # ---- 开发完成日期 / 发表状态 / 开发方式 ----
    d0, d1, d2 = _dedup(date_tbl.rows[0]), _dedup(date_tbl.rows[1]), _dedup(date_tbl.rows[2])
    set_cell(d0[1], format_date(g("开发完成日期")), value_font(d0[0]))
    if g("发表状态"):
        mark_option(d1[1], [g("发表状态")])
        if "已发表" in g("发表状态") and g("首次发表日期"):
            _fill_first_publish_date(d1[1], g("首次发表日期"))
    if g("开发方式"):
        mark_option(d2[1], [g("开发方式")])

    # ---- 著作权人（第一数据行 R1）----
    name, category, cert_type, cert_no = parse_copyright_holder(g("著作权人"))
    hr = _dedup(holder.rows[1])
    set_cell(hr[0], name, value_font(holder.rows[0].cells[0]))
    set_cell(hr[1], category, value_font(holder.rows[0].cells[1]))
    set_cell(hr[2], cert_type, value_font(holder.rows[0].cells[2]))
    set_cell(hr[3], cert_no, value_font(holder.rows[0].cells[3]))
    set_cell(hr[4], "中国", value_font(holder.rows[0].cells[4]))
    set_cell(hr[5], region_from_name(name) or g("省份/城市"), value_font(holder.rows[0].cells[5]))

    # ---- 权利说明 ----
    r0, r1 = _dedup(rights.rows[0]), _dedup(rights.rows[1])
    if g("权利取得方式"):
        mark_option(r0[1], [g("权利取得方式")], value_font(r0[0], r0[1]))
    if g("权利范围"):
        # 模板选项文字为“全部/部分”，字段值为“全部权利/部分权利”，需归一
        scope = "全部" if "全部" in g("权利范围") else "部分"
        mark_option(r1[1], [scope], value_font(r1[0], r1[1]))

    # ---- 软件功能和技术特点 ----
    rows = [_dedup(tech.rows[i]) for i in range(len(tech.rows))]
    # R0 开发的硬件环境, R1 运行的硬件环境, R2 开发操作系统, R3 开发环境/工具,
    # R4 运行平台/操作系统, R5 支撑环境/支持软件
    env_map = {
        0: g("开发的硬件环境"), 1: g("运行的硬件环境"), 2: g("开发该软件的操作系统"),
        3: g("软件开发环境 / 开发工具"), 4: g("该软件的运行平台 / 操作系统"),
        5: g("软件运行支撑环境 / 支持软件"),
    }
    for idx, val in env_map.items():
        if val:
            add_value_above_hint(rows[idx][1], val, value_font(rows[idx][0], rows[idx][1]))
    # R6 编程语言(C1) + 源程序量(C2 label, C3 值)
    if g("编程语言"):
        mark_option(rows[6][1], [g("编程语言")], value_font(rows[6][0], rows[6][1]))
    if g("源程序量"):
        set_cell(rows[6][3], g("源程序量") + " 行", value_font(rows[6][2], rows[6][3]))
    # R7 开发目的, R8 面向领域/行业, R9 主要功能
    if g("开发目的"):
        add_value_above_hint(rows[7][1], g("开发目的"), value_font(rows[7][0], rows[7][1]))
    if g("面向领域 / 行业"):
        add_value_above_hint(rows[8][1], g("面向领域 / 行业"), value_font(rows[8][0], rows[8][1]))
    if g("软件的主要功能"):
        add_value_above_hint(rows[9][1], g("软件的主要功能"), value_font(rows[9][0], rows[9][1]))
    # R10 技术类型(多选) + R11 技术特点描述
    tech_value = g("软件的技术特点")
    if tech_value:
        chosen_types = [t for t in TECH_TYPES if re.search(r"(?<![A-Za-z0-9])" + re.escape(t) + r"(?![A-Za-z0-9])", tech_value)]
        if chosen_types:
            mark_option(rows[10][1], chosen_types, value_font(rows[10][0], rows[10][1]))
        add_value_above_hint(rows[11][1], tech_value, value_font(rows[11][0], rows[11][1]))

    # ---- 申请人信息：默认由著作权人申请 ----
    ap0 = _dedup(applicant.rows[0])
    _normalize_checkbox_bullets(ap0[1])
    mark_option(ap0[1], ["由著作权人申请"])

    doc.save(str(out_path))


def default_template_path() -> Path:
    return Path(__file__).resolve().parents[1] / "assets" / "计算机软件著作权登记申请表-模板.docx"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workdir", default="软件著作权申请资料")
    ap.add_argument("--software-name", required=True)
    ap.add_argument("--version", default="V1.0")
    ap.add_argument("--template", default="", help="官方模板 docx 路径；默认用 skill 内置模板")
    args = ap.parse_args()

    if not DOCX_AVAILABLE:
        raise SystemExit("需要 python-docx：pip install python-docx")

    workdir = Path(args.workdir)
    md_path = workdir / "草稿" / "申请表信息.md"
    if not md_path.exists():
        raise SystemExit(f"STOP_FOR_USER\nNEXT_ACTION: 缺少 {md_path}，请先生成并确认申请表信息。")

    template = Path(args.template) if args.template else default_template_path()
    if not template.exists():
        raise SystemExit(f"模板不存在：{template}")

    fields = parse_application_fields(md_path)
    out_dir = ensure_dir(workdir / "正式资料")
    out_path = out_dir / "计算机软件著作权登记申请表.docx"
    build_form(args.software_name, fields, template, out_path)
    print(f"OK application form: {out_path}")


if __name__ == "__main__":
    main()
